#!/usr/bin/env python3
"""Step-through PowerPoint of the ASSUME pipeline flowchart (figures/assume_pipeline_flowchart.pptx).

One slide per build step, each slide adds the next part of the chart. Shapes are native and editable.
Uses the same step list as the flowchart animation (scripts/flowchart_spec.py).
"""

from __future__ import annotations

from pathlib import Path
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from flowchart_spec import COLORS, STEP_NAMES, steps

OUT = Path(__file__).resolve().parent.parent / "figures" / "assume_pipeline_flowchart.pptx"
S = 13.333 / 16  # inches per canvas unit
FONT = "Calibri"

NOTES = [
    "ASSUME builds an interpretable concept dictionary from unlabelled scientific images, then uses it as the embedding.",
    "DINOv3 is a frozen teacher. The sampler picks each batch: one third random, one third farthest-point in DINOv3 "
    "space, and one third collision pairs, which DINOv3 sees as different but the current concepts score the same.",
    "The proposer looks at the batch and writes new atomic concepts with 0 / 1 / 3 / 5 anchors, like a microscopist "
    "taking note that two micrographs differ in particle density.",
    "The merger prunes concepts that never vary, and for highly correlated pairs it looks at the micrographs and "
    "decides to merge them or keep both.",
    "The verifier retrieves papers for each concept through the Valency Bond MCP and removes concepts that describe "
    "how the image was taken instead of the specimen.",
    "The output is a dictionary of named rubrics, each with literature references.",
    "At run time, a new micrograph comes in without a label.",
    "The embedding agent scores it against every rubric, describing it with the vocabulary that was already built.",
    "The result is a concept vector where every dimension is a named morphological concept.",
    "The vector can be used for clustering and retrieval, and a researcher can edit one concept and search for the "
    "nearest micrograph.",
]


def rgb(h):
    return RGBColor.from_string(h)


def xy(x, y):
    """canvas (bottom-left origin) -> slide inches (top-left origin)"""
    return Inches(x * S), Inches((9 - y) * S)


def add_arrowheads(connector, head=False, tail=True):
    ln = connector.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if on:
            el = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
            el.set("type", "triangle")
            el.set("w", "med")
            el.set("h", "med")


def draw(slide, prim):
    kind = prim[0]
    sh = slide.shapes
    if kind == "box":
        _, x, y, w, h, title, sub, style = prim
        fc, ec = COLORS[style]
        left, top = xy(x, y + h)
        shp = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(w * S), Inches(h * S))
        shp.adjustments[0] = 0.12
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fc)
        shp.line.color.rgb = rgb(ec)
        shp.line.width = Pt(1.75)
        shp.shadow.inherit = False
        shp.name = title
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in ("margin_left", "margin_right"):
            setattr(tf, m, Inches(0.05))
        big = h > 1.0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        r.font.size, r.font.bold, r.font.name = Pt(15 if big else 13), True, FONT
        r.font.color.rgb = rgb(COLORS["ink"])
        for line in sub.split("\n"):
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = line
            r.font.size, r.font.name = Pt(11.5 if big else 10.5), FONT
            r.font.color.rgb = rgb(COLORS["muted"])
        tf.paragraphs[1].space_before = Pt(4)
    elif kind in ("arrow", "line"):
        pts = prim[1]
        double = prim[2] if kind == "arrow" else False
        color = COLORS["rubric"] if (kind == "arrow" and len(pts) > 2) else (COLORS["sep"] if kind == "line" else COLORS["ink"])
        for k, (a, b) in enumerate(zip(pts[:-1], pts[1:])):
            c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, *xy(*a), *xy(*b))
            c.line.color.rgb = rgb(color)
            c.line.width = Pt(2 if kind == "arrow" else 1.25)
            if kind == "line":
                c.line.dash_style = 4  # dash
            last = k == len(pts) - 2
            if kind == "arrow":
                add_arrowheads(c, head=double and k == 0, tail=last)
    elif kind == "text":
        _, x, y, s, k = prim
        size = {"title": 24, "phase": 15, "quote": 11.5, "label": 11.5}[k]
        lines = s.split("\n")
        w = {"title": 15.0, "phase": 6.5}.get(k, 4.2)
        h = 0.3 * len(lines) + 0.15
        left, top = xy(x - w / 2, y)
        tb = sh.add_textbox(left, top, Inches(w * S), Inches(h * S))
        tf = tb.text_frame
        tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = line
            r.font.size, r.font.name = Pt(size), FONT
            r.font.bold = k in ("title", "phase")
            r.font.italic = k in ("quote", "label")
            r.font.color.rgb = rgb({"title": COLORS["ink"], "phase": COLORS["ink"], "quote": COLORS["quote"],
                                    "label": COLORS["rubric"]}[k])


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    all_steps = steps()
    for k in range(len(all_steps)):
        slide = prs.slides.add_slide(blank)
        for step in all_steps[: k + 1]:
            for prim in step:
                draw(slide, prim)
        slide.notes_slide.notes_text_frame.text = f"Step {k + 1}: {STEP_NAMES[k]}\n{NOTES[k]}"
    prs.save(OUT)
    print("wrote", OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
